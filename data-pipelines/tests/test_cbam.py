import datetime
import io
from collections.abc import Callable

import openpyxl
import psycopg
import pytest

from shockt.sources import cbam, reference
from shockt.sources.cbam import CBAM_DEFAULT_VALUES, CbamSourceError, parse_workbook
from tests.conftest import MemoryArchive

COUNTRY_ROWS = [
    ("Cement", "", "", "", "", ""),
    ("2507 00 80 80", "Calcined clay", "-", "-", "-", ""),
    ("2523 10 00 90", "Other clinker", "0,870", "0,000", "0,870", "(A)"),
    ("Aluminium", "", "", "", "", ""),
    ("7601", "Unwrought aluminium", "0,360", "N/A", "0,360", "(L)"),
    ("Fertilisers", "", "", "", "", ""),
    ("3102", "Nitrogenous fertilizers", "see below", "see below", "see below", ""),
]  # fmt: skip


def workbook_bytes(
    build: Callable[[openpyxl.Workbook], None],
) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def cbam_workbook(
    rows: list[tuple[str, ...]] = COUNTRY_ROWS, header: tuple[str, ...] = cbam.HEADER
) -> bytes:
    def build(wb: openpyxl.Workbook) -> None:
        wb.active.title = "Overview"
        history = wb.create_sheet("Version History")
        history.append(["Version History"])
        history.append(["Version", "Date", "Version Notes"])
        history.append([1, datetime.datetime(2026, 2, 4), "First"])
        history.append([2, datetime.datetime(2026, 8, 6), "Corrected"])
        for title, country in (("Albania", "Albania"), ("_Other Countries", "Other Countries")):
            sheet = wb.create_sheet(title)
            sheet.append([country, "", "", "", "", ""])
            sheet.append(header)
            for row in rows:
                sheet.append(row)
        wb.create_sheet("Annex IV").append(["Annex IV"])

    return workbook_bytes(build)


def test_parse_workbook_reads_every_country_sheet_with_sectors() -> None:
    parsed = parse_workbook(cbam_workbook())

    assert parsed.rows_received == 8
    assert {row[0] for row in parsed.rows} == {"Albania", "Other Countries"}
    assert parsed.rows[1] == (
        "Albania",
        "Cement",
        "2523 10 00 90",
        "Other clinker",
        0.87,
        None,
        0.0,
        None,
        0.87,
        None,
        "(A)",
    )
    assert len(CBAM_DEFAULT_VALUES.columns) == len(parsed.rows[0]) + 1


def test_parse_workbook_keeps_the_sources_text_for_cells_that_are_not_values() -> None:
    parsed = parse_workbook(cbam_workbook())

    assert parsed.rows[0][4:] == (None, "-", None, "-", None, "-", None)
    assert parsed.rows[2][4:] == (0.36, None, None, "N/A", 0.36, None, "(L)")
    assert parsed.rows[3][4:] == (None, "see below", None, "see below", None, "see below", None)


def test_parse_workbook_reports_the_latest_version() -> None:
    parsed = parse_workbook(cbam_workbook())

    assert parsed.source_version == "version 2 (2026-08-06): Corrected"


def test_parse_workbook_rejects_an_unknown_marker() -> None:
    bad = [*COUNTRY_ROWS[:2], ("7601", "Unwrought aluminium", "tbd", "0,000", "0,000", "")]

    with pytest.raises(CbamSourceError, match="'tbd' is not a value"):
        parse_workbook(cbam_workbook(bad))


def test_parse_workbook_rejects_a_changed_header() -> None:
    header = ("CN code", *cbam.HEADER[1:])

    with pytest.raises(CbamSourceError, match="unexpected header"):
        parse_workbook(cbam_workbook(header=header))


def test_parse_workbook_rejects_a_repeated_cn_code() -> None:
    with pytest.raises(CbamSourceError, match="appears twice"):
        parse_workbook(cbam_workbook([*COUNTRY_ROWS, COUNTRY_ROWS[4]]))


def test_parse_workbook_rejects_a_row_before_any_sector_heading() -> None:
    with pytest.raises(CbamSourceError, match="before any sector heading"):
        parse_workbook(cbam_workbook(COUNTRY_ROWS[1:]))


def test_load_writes_every_row_and_archives_the_workbook(
    conn: psycopg.Connection, memory_archive: MemoryArchive
) -> None:
    body = cbam_workbook()

    result = reference.load(
        conn,
        memory_archive,
        CBAM_DEFAULT_VALUES,
        cbam.FILENAME,
        cbam.XLSX,
        {"version": "test"},
        lambda: body,
        parse_workbook,
        None,
    )

    assert result.row_count == 8
    assert memory_archive.objects["raw/cbam_default_values/latest/cbam_default_values.xlsx"] == body
    row = conn.execute(
        "SELECT direct_emissions, indirect_emissions, indirect_emissions_note"
        " FROM raw.cbam_default_values"
        " WHERE load_id = %s AND country = 'Albania' AND cn_code = '7601'",
        (result.load_id,),
    ).fetchone()
    assert row == (0.36, None, "N/A")
