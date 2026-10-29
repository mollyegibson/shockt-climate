"""Open CEDA emission factors: kgCO2e per US dollar of output by country and industry.

The workbook is archived unchanged and its `GHG_t_Raw` sheet loaded as one `latest` partition.
That sheet is a country × industry matrix of values; the other sheets only adjust it for price
year, currency and price type.
"""

import io
from collections.abc import Sequence

import openpyxl

from shockt.loading import ParsedSource, Target

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILENAME = "open_ceda.xlsx"
SHEET = "GHG_t_Raw"
FIRST_SECTOR_COLUMN = 3

OPEN_CEDA_FACTORS = Target(
    dataset="open_ceda_factors",
    table="open_ceda_factors",
    columns=(
        "load_id",
        "country_code",
        "country",
        "sector_code",
        "sector_name",
        "unit",
        "price_year",
        "price_type",
        "factor",
    ),
)


class CedaSourceError(Exception):
    """The Open CEDA workbook doesn't have the layout the parser expects."""


def parse_workbook(body: bytes) -> ParsedSource:
    """Every country × industry factor of the `GHG_t_Raw` sheet."""
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise CedaSourceError(f"Open CEDA workbook has no {SHEET!r} sheet: {workbook.sheetnames}")
    lines = list(workbook[SHEET].iter_rows(values_only=True))
    if len(lines) < 5:
        raise CedaSourceError(f"Open CEDA sheet {SHEET!r} has {len(lines)} rows")
    year_row, price_row, names, header = lines[:4]
    if year_row[0] != "Year" or price_row[0] != "Price type":
        raise CedaSourceError(f"Open CEDA sheet {SHEET!r} has no Year and Price type rows")
    if header[:2] != ("Country Code", "Country"):
        raise CedaSourceError(f"Open CEDA sheet {SHEET!r} has an unexpected header: {header[:3]}")
    price_year, price_type = int(str(year_row[1])), str(price_row[1])
    sectors = _sectors(names, header)

    rows: list[tuple[object, ...]] = []
    seen: set[str] = set()
    for line in lines[4:]:
        if not line or line[0] is None:
            continue
        code = str(line[0]).strip()
        if code in seen:
            raise CedaSourceError(f"Open CEDA country code {code} appears twice")
        seen.add(code)
        unit = str(line[2]).strip() if line[2] is not None else ""
        if not unit or line[1] is None:
            raise CedaSourceError(f"Open CEDA country {code} has no name or unit")
        for position, (sector_code, sector_name) in sectors.items():
            value = line[position]
            if value is not None and not isinstance(value, int | float):
                raise CedaSourceError(f"Open CEDA {code} {sector_code} is not a number: {value!r}")
            rows.append(
                (
                    code,
                    str(line[1]).strip(),
                    sector_code,
                    sector_name,
                    unit,
                    price_year,
                    price_type,
                    None if value is None else float(value),
                )
            )
    if not rows:
        raise CedaSourceError(f"Open CEDA sheet {SHEET!r} has no country rows")
    return ParsedSource(rows=rows, rows_received=len(rows), source_version=None)


def _sectors(names: Sequence[object], header: Sequence[object]) -> dict[int, tuple[str, str]]:
    """Column position → (sector code, sector name) for every industry column."""
    sectors: dict[int, tuple[str, str]] = {}
    for position in range(FIRST_SECTOR_COLUMN, len(header)):
        if header[position] is None:
            continue
        code = str(header[position]).strip()
        name = names[position] if position < len(names) else None
        if name is None:
            raise CedaSourceError(f"Open CEDA sector {code} has no name")
        if any(code == existing[0] for existing in sectors.values()):
            raise CedaSourceError(f"Open CEDA sector code {code} appears twice")
        sectors[position] = (code, str(name).strip())
    if not sectors:
        raise CedaSourceError(f"Open CEDA sheet {SHEET!r} has no industry columns")
    return sectors
