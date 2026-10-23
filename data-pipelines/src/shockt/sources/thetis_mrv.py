"""EMSA THETIS-MRV: annual per-ship fuel consumption and emissions under EU MRV.

Each reporting year's xlsx is downloaded from the THETIS-MRV public API, archived unchanged,
then parsed with openpyxl. The 2024+ file has 113 columns (CH₄/N₂O/CO₂eq added) and two sheets
("Full ERs" and "Partial ERs"); pre-2024 has 62 columns and one sheet. The parser finds columns
by name from the header row, so it handles both formats.
"""

import io
from collections.abc import Callable
from datetime import date

import httpx
import openpyxl
import psycopg

from shockt import http
from shockt.archive import Archive, object_key
from shockt.config import ThetisMrvConfig
from shockt.loading import LoadResult, ParsedSource, Target, load_partition

DATASET = "thetis_mrv"

TARGET = Target(
    dataset=DATASET,
    table="thetis_mrv",
    columns=(
        "load_id",
        "reporting_year",
        "report_type",
        "imo_number",
        "ship_name",
        "ship_type",
        "port_of_registry",
        "home_port",
        "ice_class",
        "doc_issue_date",
        "doc_expiry_date",
        "total_fuel_consumption_mt",
        "total_co2_mt",
        "total_co2eq_mt",
        "time_at_sea_hours",
        "fuel_per_distance_kg_nm",
        "co2_per_distance_kg_nm",
        "co2_per_transport_work_mass",
    ),
)

HEADER_ROW = 3

REQUIRED_COLUMNS = {
    "IMO Number",
    "Name",
    "Ship type",
    "Home Port",
    "Ice Class",
    "DoC issue date",
    "DoC expiry date",
    "Reporting Period",
    "Total fuel consumption [m tonnes]",
    "Total CO₂ emissions [m tonnes]",
}

OPTIONAL_COLUMNS = {
    "Total CO₂eq emissions [m tonnes]",
    "Annual average Fuel consumption per distance [kg / n mile]",
    "Fuel consumption per distance [kg / n mile]",
    "Annual average CO₂ emissions per distance [kg CO₂ / n mile]",
    "CO₂ emissions per distance [kg CO₂ / n mile]",
    "Annual average CO₂ emissions per transport work (mass) [g CO₂ / m tonnes · n miles]",
    "CO₂ emissions per transport work (mass) [g CO₂ / m tonnes · n miles]",
}


class MrvSourceError(Exception):
    """The MRV xlsx doesn't have the layout the parser expects."""


def _safe_float(value: object) -> float | None:
    """Convert a cell value to float, treating MRV's error strings as null."""
    if value is None or value == "N/A" or value == "Division by zero!":
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped or stripped == "N/A" or stripped == "Division by zero!":
            return None
        return float(stripped)
    return float(value)


def _safe_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    return None


def _find_column(headers: list[object], *candidates: str) -> int | None:
    """Find the first matching column name from several candidates."""
    for candidate in candidates:
        try:
            return headers.index(candidate)
        except ValueError:
            continue
    return None


def _parse_sheet(
    workbook: openpyxl.Workbook,
    sheet_name: str,
    reporting_year: int,
    report_type: str,
) -> list[tuple[object, ...]]:
    """Parse one sheet of an MRV xlsx into rows for raw.thetis_mrv."""
    ws = workbook[sheet_name]
    headers: list[object] = []
    for row in ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True):
        headers = list(row)
    if not headers:
        raise MrvSourceError(f"sheet {sheet_name!r} has no header at row {HEADER_ROW}")

    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise MrvSourceError(f"sheet {sheet_name!r} missing required columns: {missing}")

    imo_col = headers.index("IMO Number")
    name_col = headers.index("Name")
    type_col = headers.index("Ship type")
    flag_col = _find_column(headers, "Flag State", "Port of Registry")
    home_col = headers.index("Home Port")
    ice_col = headers.index("Ice Class")
    doc_issue_col = headers.index("DoC issue date")
    doc_expiry_col = headers.index("DoC expiry date")
    fuel_col = headers.index("Total fuel consumption [m tonnes]")
    co2_col = headers.index("Total CO₂ emissions [m tonnes]")
    time_col = _find_column(
        headers,
        "Time spent at sea [hours]",
        "Annual Time spent at sea [hours]",
    )

    co2eq_col = _find_column(headers, "Total CO₂eq emissions [m tonnes]")
    fuel_dist_col = _find_column(
        headers,
        "Annual average Fuel consumption per distance [kg / n mile]",
        "Fuel consumption per distance [kg / n mile]",
    )
    co2_dist_col = _find_column(
        headers,
        "Annual average CO₂ emissions per distance [kg CO₂ / n mile]",
        "CO₂ emissions per distance [kg CO₂ / n mile]",
    )
    co2_tw_col = _find_column(
        headers,
        "Annual average CO₂ emissions per transport work (mass) [g CO₂ / m tonnes · n miles]",
        "CO₂ emissions per transport work (mass) [g CO₂ / m tonnes · n miles]",
    )

    rows: list[tuple[object, ...]] = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        cells = list(row)
        if cells[imo_col] is None:
            continue
        rows.append(
            (
                reporting_year,
                report_type,
                str(cells[imo_col]).strip(),
                cells[name_col] if cells[name_col] is not None else None,
                cells[type_col] if cells[type_col] is not None else None,
                cells[flag_col] if flag_col is not None else None,
                cells[home_col] if cells[home_col] is not None else None,
                cells[ice_col] if cells[ice_col] is not None else None,
                _safe_date(cells[doc_issue_col]),
                _safe_date(cells[doc_expiry_col]),
                _safe_float(cells[fuel_col]),
                _safe_float(cells[co2_col]),
                _safe_float(cells[co2eq_col]) if co2eq_col is not None else None,
                _safe_float(cells[time_col]) if time_col is not None else None,
                _safe_float(cells[fuel_dist_col]) if fuel_dist_col is not None else None,
                _safe_float(cells[co2_dist_col]) if co2_dist_col is not None else None,
                _safe_float(cells[co2_tw_col]) if co2_tw_col is not None else None,
            )
        )
    return rows


def available_versions(client: httpx.Client, config: ThetisMrvConfig) -> dict[int, int]:
    """Reporting year → latest version number, from the downloadable-files endpoint."""
    response = http.get(client, config.downloadable_files_url, {})
    body = response.json()
    entries = body.get("results", body) if isinstance(body, dict) else body
    versions: dict[int, int] = {}
    for entry in entries:
        year = entry.get("reportingPeriod") or entry.get("year")
        version = entry.get("version")
        if year is not None and version is not None:
            year_int = int(year)
            version_int = int(version)
            if year_int not in versions or version_int > versions[year_int]:
                versions[year_int] = version_int
    return versions


def fetch_year(client: httpx.Client, config: ThetisMrvConfig, year: int, version: int) -> bytes:
    """Download one reporting year's xlsx."""
    url = config.download_url_template.format(year=year, version=version)
    response = http.get(client, url, {})
    return response.content


def parse_xlsx(body: bytes, reporting_year: int) -> ParsedSource:
    """Parse one year's MRV xlsx (pre-2024 or 2024+ format) into rows for raw.thetis_mrv."""
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    all_rows: list[tuple[object, ...]] = []

    full_sheet = None
    partial_sheet = None
    for name in workbook.sheetnames:
        lower = name.lower()
        if "partial" in lower:
            partial_sheet = name
        elif "full" in lower or len(workbook.sheetnames) == 1:
            full_sheet = name

    if full_sheet is None and len(workbook.sheetnames) == 1:
        full_sheet = workbook.sheetnames[0]

    if full_sheet is None:
        raise MrvSourceError(f"no 'Full ERs' sheet found in {workbook.sheetnames}")

    all_rows.extend(_parse_sheet(workbook, full_sheet, reporting_year, "full"))
    if partial_sheet is not None:
        all_rows.extend(_parse_sheet(workbook, partial_sheet, reporting_year, "partial"))

    workbook.close()
    return ParsedSource(
        rows=all_rows,
        rows_received=len(all_rows),
        source_version=None,
    )


def load_year(
    conn: psycopg.Connection,
    archive: Archive,
    config: ThetisMrvConfig,
    year: int,
    version: int,
    obtain: Callable[[], bytes] | None,
) -> LoadResult:
    """Archive the xlsx, parse it and load one reporting year."""
    key = object_key(DATASET, str(year), f"thetis_mrv_{year}.xlsx")
    params: dict[str, object] = {
        "reporting_year": year,
        "version": version,
    }
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return load_partition(
        conn,
        archive,
        TARGET,
        str(year),
        key=key,
        content_type=xlsx_type,
        params=params,
        obtain=obtain,
        parse=lambda body: parse_xlsx(body, year),
    )
