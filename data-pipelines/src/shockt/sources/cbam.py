"""EU CBAM default values for the definitive period (Annex I), loaded whole.

The workbook has one sheet per exporting country, each a table of CN codes grouped under sector
headings. The file is archived unchanged and loaded as one `latest` partition. Values keep the
source's units (tCO2eq per tonne of good).
"""

import io
import re
from collections.abc import Iterable

import openpyxl

from shockt.loading import ParsedSource, Target

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILENAME = "cbam_default_values.xlsx"
NON_COUNTRY_SHEETS = frozenset({"Overview", "Version History", "Annex IV"})
HEADER = (
    "Product CN Code / TARIC Code",
    "Description",
    "Default Value (direct emissions) (tCO2eq/tonne of good)",
    "Default Value (indirect emissions) (tCO2eq/tonne of good)",
    "Default Value (total emissions) (tCO2eq/tonne of good)",
    "Underlying production route determining CBAM BM",
)
NOT_A_VALUE = frozenset({"-", "N/A", "see below"})
DECIMAL_COMMA = re.compile(r"-?\d+(,\d+)?")

CBAM_DEFAULT_VALUES = Target(
    dataset="cbam_default_values",
    table="cbam_default_values",
    columns=(
        "load_id",
        "country",
        "sector",
        "cn_code",
        "description",
        "direct_emissions",
        "direct_emissions_note",
        "indirect_emissions",
        "indirect_emissions_note",
        "total_emissions",
        "total_emissions_note",
        "production_route",
    ),
)


class CbamSourceError(Exception):
    """The CBAM workbook doesn't have the layout the parser expects."""


def parse_workbook(body: bytes) -> ParsedSource:
    """Every default value of every country sheet, with the workbook's latest version."""
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    countries = [n for n in workbook.sheetnames if n not in NON_COUNTRY_SHEETS]
    if not countries:
        raise CbamSourceError(f"CBAM workbook has no country sheets: {workbook.sheetnames}")
    rows: list[tuple[object, ...]] = []
    seen: set[tuple[str, str]] = set()
    for name in countries:
        for row in _sheet_rows(name, workbook[name].iter_rows(values_only=True)):
            key = (str(row[0]), str(row[2]))
            if key in seen:
                raise CbamSourceError(f"CBAM sheet {name!r}: CN code {row[2]} appears twice")
            seen.add(key)
            rows.append(row)
    return ParsedSource(
        rows=rows,
        rows_received=len(rows),
        source_version=_version(workbook["Version History"].iter_rows(values_only=True)),
    )


def _sheet_rows(title: str, rows: Iterable[tuple[object, ...]]) -> list[tuple[object, ...]]:
    lines = [_padded(row) for row in rows]
    if len(lines) < 3 or not _text(lines[0][0]):
        raise CbamSourceError(f"CBAM sheet {title!r} has no country title")
    country = str(lines[0][0]).strip()
    if lines[1] != HEADER:
        raise CbamSourceError(f"CBAM sheet {title!r} has an unexpected header: {lines[1]}")
    parsed: list[tuple[object, ...]] = []
    sector: str | None = None
    for line in lines[2:]:
        if not any(line):
            continue
        if _text(line[0]) and not any(_text(c) for c in line[1:]):
            sector = str(line[0]).strip()
            continue
        if sector is None:
            raise CbamSourceError(f"CBAM sheet {title!r}: row before any sector heading: {line}")
        direct, direct_note = _value(line[2], title)
        indirect, indirect_note = _value(line[3], title)
        total, total_note = _value(line[4], title)
        parsed.append(
            (
                country,
                sector,
                str(line[0]).strip(),
                _text(line[1]),
                direct,
                direct_note,
                indirect,
                indirect_note,
                total,
                total_note,
                _text(line[5]),
            )
        )
    return parsed


def _version(rows: Iterable[tuple[object, ...]]) -> str:
    """The last row of the Version History sheet: version number, date and notes."""
    entries = [row for row in rows if isinstance(row[0], int) and row[1]]
    if not entries:
        raise CbamSourceError("CBAM workbook has no Version History entries")
    number, published, notes = entries[-1][0], entries[-1][1], entries[-1][2]
    return f"version {number} ({published:%Y-%m-%d}): {notes}"


def _value(cell: object, sheet: str) -> tuple[float | None, str | None]:
    """A default value as a number, or null with the source's text when the cell isn't one."""
    if cell is None or cell == "":
        return None, None
    if isinstance(cell, int | float):
        return float(cell), None
    text = str(cell).strip()
    if text in NOT_A_VALUE:
        return None, text
    if DECIMAL_COMMA.fullmatch(text):
        return float(text.replace(",", ".")), None
    raise CbamSourceError(f"CBAM sheet {sheet!r}: {text!r} is not a value or a known marker")


def _padded(row: tuple[object, ...]) -> tuple[object, ...]:
    return (*row, *([None] * (len(HEADER) - len(row))))[: len(HEADER)]


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
