"""Freight emission factor sources, loaded whole: the UK government GHG conversion factors flat
file, and Tables 8 (transportation) and 11 (GWPs) of the EPA GHG Emission Factors Hub.

Each workbook is archived unchanged and loaded as one `latest` partition. Values keep the
source's labels and units; `core.ref_emission_factors` derives the freight factors from them.
"""

import io
from collections.abc import Sequence

import openpyxl

from shockt.loading import ParsedSource, Target

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
UK_SHEET = "Factors by Category"
UK_COLUMNS = (
    "ID",
    "Scope",
    "Level 1",
    "Level 2",
    "Level 3",
    "Level 4",
    "Column Text",
    "UOM",
    "GHG/Unit",
)
UK_VALUE_PREFIX = "GHG Conversion Factor"

UK_FACTORS = Target(
    dataset="uk_ghg_conversion_factors",
    table="uk_ghg_conversion_factors",
    columns=(
        "load_id",
        "id",
        "scope",
        "level_1",
        "level_2",
        "level_3",
        "level_4",
        "column_text",
        "uom",
        "ghg_unit",
        "factor",
    ),
)
EPA_TRANSPORT = Target(
    dataset="epa_ghg_transport_factors",
    table="epa_ghg_transport_factors",
    columns=("load_id", "vehicle_type", "co2_factor", "ch4_factor", "n2o_factor", "units"),
)
EPA_GWP = Target(
    dataset="epa_ghg_gwp",
    table="epa_ghg_gwp",
    columns=("load_id", "common_name", "chemical_formula", "gwp_100yr"),
)


class FactorSourceError(Exception):
    """A factor workbook doesn't have the layout the parser expects."""


def parse_uk_workbook(body: bytes) -> ParsedSource:
    """Every factor row of the UK flat file's "Factors by Category" sheet."""
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if UK_SHEET not in workbook.sheetnames:
        raise FactorSourceError(f"UK workbook has no {UK_SHEET!r} sheet: {workbook.sheetnames}")
    rows = list(workbook[UK_SHEET].iter_rows(values_only=True))
    header_index = next((i for i, row in enumerate(rows) if row and "Level 1" in row), None)
    if header_index is None:
        raise FactorSourceError(f"UK sheet {UK_SHEET!r} has no header row")
    header = rows[header_index]
    missing = [c for c in UK_COLUMNS if c not in header]
    values = [c for c in header if isinstance(c, str) and c.startswith(UK_VALUE_PREFIX)]
    if missing or len(values) != 1:
        raise FactorSourceError(
            f"UK header is missing {missing} or has {len(values)} {UK_VALUE_PREFIX!r} columns"
        )
    positions = [header.index(c) for c in (*UK_COLUMNS, values[0])]
    parsed: list[tuple[object, ...]] = []
    seen: set[str] = set()
    for row in rows[header_index + 1 :]:
        if not row or row[positions[0]] is None:
            continue
        record = [row[p] for p in positions]
        factor_id = str(record[0])
        if factor_id in seen:
            raise FactorSourceError(f"UK factor ID {factor_id} appears twice")
        seen.add(factor_id)
        factor = record[-1]
        if factor is not None and not isinstance(factor, int | float):
            raise FactorSourceError(f"UK factor {factor_id} is not a number: {factor!r}")
        parsed.append(
            (
                factor_id,
                *(None if v is None else str(v) for v in record[1:-1]),
                None if factor is None else float(factor),
            )
        )
    return ParsedSource(rows=parsed, rows_received=len(parsed), source_version=None)


def parse_epa_transport(body: bytes) -> ParsedSource:
    """Every row of EPA Hub Table 8: kg CO2, g CH4 and g N2O per vehicle-mile or ton-mile."""
    rows = _table_rows(_epa_rows(body), "Table 8", "Vehicle Type")
    parsed: list[tuple[object, ...]] = []
    for row in rows:
        if len(row) != 5 or not all(isinstance(v, int | float) for v in row[1:4]):
            raise FactorSourceError(f"EPA Table 8 row has an unexpected shape: {row}")
        parsed.append((str(row[0]), *(float(v) for v in row[1:4]), str(row[4])))
    return ParsedSource(rows=parsed, rows_received=len(parsed), source_version=None)


def parse_epa_gwp(body: bytes) -> ParsedSource:
    """Every gas in EPA Hub Table 11 with its 100-year GWP."""
    rows = _table_rows(_epa_rows(body), "Table 11", "Industrial Designation")
    parsed: list[tuple[object, ...]] = []
    for row in rows:
        if len(row) != 3 or not isinstance(row[2], int | float):
            raise FactorSourceError(f"EPA Table 11 row has an unexpected shape: {row}")
        parsed.append((str(row[0]), str(row[1]), float(row[2])))
    return ParsedSource(rows=parsed, rows_received=len(parsed), source_version=None)


def _epa_rows(body: bytes) -> list[list[object]]:
    sheet = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True).worksheets[0]
    return [[c for c in row if c is not None] for row in sheet.iter_rows(values_only=True)]


def _table_rows(rows: Sequence[list[object]], title: str, header_start: str) -> list[list[object]]:
    """The data rows of one EPA table: after its header row, up to its source note or the next
    table."""
    start = next((i for i, row in enumerate(rows) if row and row[0] == title), None)
    if start is None:
        raise FactorSourceError(f"EPA workbook has no {title}")
    header = next(
        (
            i
            for i in range(start + 1, len(rows))
            if rows[i] and str(rows[i][0]).startswith(header_start)
        ),
        None,
    )
    if header is None:
        raise FactorSourceError(f"EPA {title} has no header row starting {header_start!r}")
    data: list[list[object]] = []
    for row in rows[header + 1 :]:
        first = str(row[0]) if row else ""
        if first.startswith(("Source", "Notes", "Table ")):
            break
        if row:
            data.append(row)
    if not data:
        raise FactorSourceError(f"EPA {title} has no rows")
    return data
